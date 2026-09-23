"""Document operations for office_tool."""


def _run_document_op(op: str, args: dict) -> str:
    """Eksekusi operasi document menggunakan python-docx."""
    input_path = args.get("input", "")
    output_path = args.get("output", "")

    if op == "create_docx":
        try:
            from docx import Document

            doc = Document()
            title = args.get("title", "")
            content = args.get("content", "")
            paragraphs = args.get("paragraphs", [])
            if title:
                doc.add_heading(title, 0)
            if content:
                doc.add_paragraph(content)
            for p in paragraphs:
                if isinstance(p, dict):
                    style = p.get("style", "Normal")
                    text = p.get("text", "")
                    doc.add_paragraph(text, style=style)
                else:
                    doc.add_paragraph(str(p))
            doc.save(output_path)
            return f"OK: Word dibuat -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    elif op == "edit_docx":
        try:
            from docx import Document

            doc = Document(input_path)
            action = args.get("action", "append")
            if action == "append":
                content = args.get("content", "")
                paragraphs = args.get("paragraphs", [])
                if content:
                    doc.add_paragraph(content)
                for p in paragraphs:
                    doc.add_paragraph(str(p))
            elif action == "replace":
                old_text = args.get("old_text", "")
                new_text = args.get("new_text", "")
                for p in doc.paragraphs:
                    if old_text in p.text:
                        p.text = p.text.replace(old_text, new_text)
            doc.save(output_path)
            return f"OK: Word diedit -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    elif op == "merge_docx":
        try:
            from docx import Document

            inputs = args.get("inputs", [input_path])
            merged = Document()
            for f in inputs:
                doc = Document(f)
                for p in doc.paragraphs:
                    merged.add_paragraph(p.text, style=p.style)
            merged.save(output_path)
            return f"OK: {len(inputs)} Word digabung -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    elif op == "template_fill":
        try:
            from docx import Document

            doc = Document(input_path)
            replacements = args.get("replacements", {})
            for p in doc.paragraphs:
                for key, value in replacements.items():
                    if key in p.text:
                        p.text = p.text.replace(key, str(value))
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for key, value in replacements.items():
                            if key in cell.text:
                                cell.text = cell.text.replace(key, str(value))
            doc.save(output_path)
            return f"OK: Template diisi ({len(replacements)} field) -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    elif op == "docx_to_md":
        try:
            from docx import Document

            doc = Document(input_path)
            md_lines = []
            for p in doc.paragraphs:
                style = p.style
                if style and style.name.startswith("Heading"):
                    level = int(style.name[-1]) if style.name[-1].isdigit() else 1
                    md_lines.append(f"{'#' * level} {p.text}")
                elif style and style.name == "List Bullet":
                    md_lines.append(f"- {p.text}")
                elif style and style.name == "List Number":
                    md_lines.append(f"1. {p.text}")
                else:
                    md_lines.append(p.text)
            with open(output_path, "w", encoding="utf-8") as f:
                f.write("\n\n".join(md_lines))
            return f"OK: Word -> Markdown -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    elif op == "md_to_docx":
        try:
            from docx import Document

            with open(input_path, encoding="utf-8") as f:
                content = f.read()
            doc = Document()
            lines = content.split("\n")
            for raw_line in lines:
                stripped = raw_line.strip()
                if not stripped:
                    doc.add_paragraph("")
                elif stripped.startswith("# "):
                    doc.add_heading(stripped[2:], 1)
                elif stripped.startswith("## "):
                    doc.add_heading(stripped[3:], 2)
                elif stripped.startswith("### "):
                    doc.add_heading(stripped[4:], 3)
                elif stripped.startswith("- "):
                    doc.add_paragraph(stripped[2:], style="List Bullet")
                elif stripped.startswith("1. "):
                    doc.add_paragraph(stripped[3:], style="List Number")
                else:
                    doc.add_paragraph(stripped)
            doc.save(output_path)
            return f"OK: Markdown -> Word -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    elif op == "html_to_docx":
        try:
            from html.parser import HTMLParser

            from docx import Document

            with open(input_path, encoding="utf-8") as f:
                html_content = f.read()
            doc = Document()

            class SimpleHTMLParser(HTMLParser):
                def __init__(self, doc):
                    super().__init__()
                    self.doc = doc
                    self.current_text = ""

                def handle_data(self, data):
                    self.current_text += data

                def handle_endtag(self, tag):
                    if tag in ("p", "div", "br"):
                        if self.current_text.strip():
                            self.doc.add_paragraph(self.current_text.strip())
                        self.current_text = ""

            parser = SimpleHTMLParser(doc)
            parser.feed(html_content)
            if parser.current_text.strip():
                doc.add_paragraph(parser.current_text.strip())
            doc.save(output_path)
            return f"OK: HTML -> Word -> {output_path}"
        except ImportError:
            return "Error: python-docx belum install. Jalankan: pip install python-docx"

    else:
        return f"Error: Operasi Document '{op}' tidak dikenal."
