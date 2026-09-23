"""Spreadsheet operations for office_tool."""

import json as _json


def _run_spreadsheet_op(op: str, args: dict) -> str:
    """Eksekusi operasi spreadsheet menggunakan openpyxl."""
    input_path = args.get("input", "")
    output_path = args.get("output", "")

    if op == "read_xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(input_path, data_only=True)
            try:
                result = {}
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        rows.append([str(c) if c is not None else "" for c in row])
                    result[sheet_name] = rows
            finally:
                wb.close()
            return _json.dumps(result, indent=2, ensure_ascii=False)
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "create_xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill

            data = args.get("data", [])
            headers = args.get("headers", [])
            title = args.get("title", "Sheet1")
            wb = Workbook()
            ws = wb.active
            ws.title = title
            if headers:
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="4472C4")
                    cell.alignment = Alignment(horizontal="center")
            for row in data:
                ws.append(row)
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)
            wb.save(output_path)
            return f"OK: Excel dibuat -> {output_path} ({len(data)} baris)"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "apply_formula":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(input_path)
            try:
                ws = wb.active
                cell = args.get("cell", "D1")
                formula = args.get("formula", "=SUM(A1:C1)")
                ws[cell] = formula
                wb.save(output_path)
            finally:
                wb.close()
            return f"OK: Formula '{formula}' diapply ke {cell} -> {output_path}"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "create_chart":
        try:
            from openpyxl import load_workbook
            from openpyxl.chart import BarChart, LineChart, PieChart, Reference

            wb = load_workbook(input_path)
            try:
                ws = wb.active
                chart_type = args.get("type", "bar")
                title = args.get("title", "Chart")
                data_ref = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=min(ws.max_row, 10))
                cats_ref = Reference(ws, min_col=1, min_row=2, max_row=min(ws.max_row, 10))
                chart = {"bar": BarChart, "pie": PieChart, "line": LineChart}.get(chart_type, BarChart)()
                chart.title = title
                chart.add_data(data_ref, titles_from_data=True)
                chart.set_categories(cats_ref)
                pos = args.get("position", "E1")
                ws.add_chart(chart, pos)
                wb.save(output_path)
            finally:
                wb.close()
            return f"OK: Chart {chart_type} dibuat di {pos} -> {output_path}"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "analyze_data":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(input_path, data_only=True)
            try:
                ws = wb.active
                stats = {"total_rows": ws.max_row, "total_cols": ws.max_column, "sheets": wb.sheetnames}
                numeric_cols = []
                for col in range(1, ws.max_column + 1):
                    values = []
                    for row in range(2, min(ws.max_row + 1, 100)):
                        v = ws.cell(row=row, column=col).value
                        if isinstance(v, (int, float)):
                            values.append(v)
                    if values:
                        col_letter = ws.cell(row=1, column=col).value or f"Col{col}"
                        numeric_cols.append(
                            {
                                "column": str(col_letter),
                                "count": len(values),
                                "sum": round(sum(values), 2),
                                "avg": round(sum(values) / len(values), 2),
                                "min": min(values),
                                "max": max(values),
                            }
                        )
                stats["numeric_columns"] = numeric_cols
            finally:
                wb.close()
            return _json.dumps(stats, indent=2, ensure_ascii=False)
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "format_cells":
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Alignment, Font, PatternFill

            wb = load_workbook(input_path)
            try:
                ws = wb.active
                cell_range = args.get("range", "A1")
                bold = args.get("bold", False)
                bg_color = args.get("bg_color")
                font_color = args.get("font_color")
                for row in ws[cell_range]:
                    for cell in row:
                        if bold:
                            cell.font = Font(bold=True)
                        if bg_color:
                            cell.fill = PatternFill("solid", fgColor=bg_color)
                        if font_color:
                            cell.font = Font(color=font_color)
                wb.save(output_path)
            finally:
                wb.close()
            return f"OK: Format diapply ke {cell_range} -> {output_path}"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "xlsx_to_csv":
        try:
            import csv

            from openpyxl import load_workbook

            wb = load_workbook(input_path, data_only=True)
            try:
                ws = wb.active
                with open(output_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    for row in ws.iter_rows(values_only=True):
                        writer.writerow([str(c) if c is not None else "" for c in row])
            finally:
                wb.close()
            return f"OK: XLSX -> CSV -> {output_path}"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "csv_to_xlsx":
        try:
            import csv

            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            with open(input_path, encoding="utf-8") as f:
                reader = csv.reader(f)
                for row in reader:
                    ws.append(row)
            wb.save(output_path)
            return f"OK: CSV -> XLSX -> {output_path}"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "json_to_xlsx":
        try:
            from openpyxl import Workbook

            data = args.get("data", [])
            wb = Workbook()
            ws = wb.active
            if data:
                headers = list(data[0].keys())
                ws.append(headers)
                for row in data:
                    ws.append([row.get(h, "") for h in headers])
            wb.save(output_path)
            return f"OK: JSON -> XLSX -> {output_path} ({len(data)} baris)"
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    elif op == "xlsx_to_json":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(input_path, data_only=True)
            try:
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    return "[]"
                headers = [str(h) for h in rows[0]]
                data = []
                for row in rows[1:]:
                    data.append(dict(zip(headers, [str(c) if c is not None else "" for c in row], strict=False)))
            finally:
                wb.close()
            return _json.dumps(data, indent=2, ensure_ascii=False)
        except ImportError:
            return "Error: openpyxl belum install. Jalankan: pip install openpyxl"

    else:
        return f"Error: Operasi Spreadsheet '{op}' tidak dikenal."
