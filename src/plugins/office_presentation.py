"""Presentation operations for office_tool."""

import json as _json
from typing import Any


def _run_presentation_op(op: str, args: dict) -> str:
    """Eksekusi operasi presentation menggunakan python-pptx."""
    input_path = args.get("input", "")
    output_path = args.get("output", "")

    if op == "create_pptx":
        try:
            from pptx import Presentation

            prs = Presentation()
            slides_data = args.get("slides", [])
            title = args.get("title", "")
            if title and not slides_data:
                slides_data = [{"title": title, "content": args.get("content", "")}]
            for slide_info in slides_data:
                layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = slide_info.get("title", "Slide")
                body = slide.placeholders[1]
                content = slide_info.get("content", "")
                if isinstance(content, list):
                    body.text = "\n".join(content)
                else:
                    body.text = str(content)
            prs.save(output_path)
            return f"OK: PowerPoint dibuat -> {output_path} ({len(slides_data)} slides)"
        except ImportError:
            return "Error: python-pptx belum install. Jalankan: pip install python-pptx"

    elif op == "extract_pptx":
        try:
            from pptx import Presentation

            prs = Presentation(input_path)
            extracted: list[dict[str, Any]] = []
            for i, slide in enumerate(prs.slides):
                slide_data: dict[str, Any] = {"slide": i + 1, "texts": []}
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_data["texts"].append(shape.text.strip())
                extracted.append(slide_data)
            return _json.dumps(extracted, indent=2, ensure_ascii=False)
        except ImportError:
            return "Error: python-pptx belum install. Jalankan: pip install python-pptx"

    elif op == "md_to_pptx":
        try:
            from pptx import Presentation

            with open(input_path, encoding="utf-8") as f:
                content = f.read()
            prs = Presentation()
            slides = content.split("\n---\n")
            for slide_content in slides:
                lines = [line.strip() for line in slide_content.strip().split("\n") if line.strip()]
                if not lines:
                    continue
                title = lines[0].lstrip("# ").strip()
                body_text = "\n".join(lines[1:])
                layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = title
                slide.placeholders[1].text = body_text
            prs.save(output_path)
            return f"OK: Markdown -> PowerPoint -> {output_path} ({len(slides)} slides)"
        except ImportError:
            return "Error: python-pptx belum install. Jalankan: pip install python-pptx"

    elif op == "md_to_slides":
        return _run_presentation_op("md_to_pptx", args)

    elif op == "add_notes":
        try:
            from pptx import Presentation

            prs = Presentation(input_path)
            slide_num = args.get("slide", 1) - 1
            notes = args.get("notes", "")
            if 0 <= slide_num < len(prs.slides):
                slide = prs.slides[slide_num]
                if not slide.has_notes_slide:
                    _ = slide.notes_slide
                notes_slide = slide.notes_slide
                notes_slide.notes_text_frame.text = notes
                prs.save(output_path)
                return f"OK: Catatan ditambahkan ke slide {slide_num + 1} -> {output_path}"
            return f"Error: Slide {slide_num + 1} tidak ditemukan"
        except ImportError:
            return "Error: python-pptx belum install. Jalankan: pip install python-pptx"

    elif op == "export_pptx":
        try:
            from pptx import Presentation

            prs = Presentation(input_path)
            export_format = args.get("format", "images")
            if export_format == "images":
                import os

                os.makedirs(output_path, exist_ok=True)
                for i, slide in enumerate(prs.slides):
                    texts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            texts.append(shape.text.strip())
                    with open(f"{output_path}/slide_{i + 1}.txt", "w", encoding="utf-8") as f:
                        f.write("\n".join(texts))
                return f"OK: Slides diekspor ke {output_path}"
            return f"Error: Format export '{export_format}' tidak didukung"
        except ImportError:
            return "Error: python-pptx belum install. Jalankan: pip install python-pptx"

    else:
        return f"Error: Operasi Presentation '{op}' tidak dikenal."
